"""
Excel-модель v2 с живыми формулами:  python v2/model/build_excel_v2.py
→ v2/model/Ucell_v2_ARPU_Loyalty_Model.xlsx

Листы: README · Inputs · Initiatives · Monthly_I · Loyalty_Inputs · Prizes · Daily_Drops ·
Loyalty_Monthly · MVP_Budget · Thresholds · Portfolio · AB_Calculator · MC_Results.
Детерминированный расчёт на модальных значениях совпадает с model_v2.py (det=True).
"""
from __future__ import annotations

import json
import os
import sys

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as CL

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import assumptions_v2 as A  # noqa: E402

OUT = os.path.join(HERE, "Ucell_v2_ARPU_Loyalty_Model.xlsx")
H = A.HORIZON_MONTHS
FONT = "Arial"
BLUE = Font(name=FONT, color="0000FF")
BLACK = Font(name=FONT, color="000000")
GREEN = Font(name=FONT, color="008000")
BOLD = Font(name=FONT, bold=True)
TITLE = Font(name=FONT, bold=True, size=14)
HDR_FILL = PatternFill("solid", start_color="1F3864")
HDR_FONT = Font(name=FONT, bold=True, color="FFFFFF")
YELLOW = PatternFill("solid", start_color="FFFF00")
SEC = PatternFill("solid", start_color="D9E1F2")
thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
NUM = '#,##0;(#,##0);"-"'
NUM1 = '#,##0.0;(#,##0.0);"-"'
NUM2 = '#,##0.00;(#,##0.00);"-"'
PCT = '0.0%;(0.0%);"-"'
PCT2 = '0.00%;(0.00%);"-"'
MULT = '0.0"x";(0.0"x");"-"'
M0 = 3  # колонка C = месяц 1


def mode(v):
    return v[1] if isinstance(v, (tuple, list)) else v


def mc(m):
    return CL(M0 + m - 1)


def hdr(ws, row, ncol, start=1):
    for c in range(start, ncol + 1):
        x = ws.cell(row=row, column=c)
        x.fill, x.font, x.border = HDR_FILL, HDR_FONT, BORDER
        x.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def put(ws, r, c, v, font=BLACK, fmt=None, fill=None, comment=None):
    x = ws.cell(row=r, column=c, value=v)
    x.font, x.border = font, BORDER
    if fmt:
        x.number_format = fmt
    if fill:
        x.fill = fill
    if comment:
        x.comment = Comment(comment, "model")
    return x


def section(ws, r, text, ncol):
    ws.cell(row=r, column=1, value=text).font = BOLD
    for c in range(1, ncol + 1):
        ws.cell(row=r, column=c).fill = SEC


# ============================================================================ README
def sheet_readme(wb):
    ws = wb.active
    ws.title = "README"
    lines = [
        ("Ucell v2 — 4 инициативы роста ARPU архивной базы + Loyalty program", TITLE),
        ("", None),
        ("I1 Sunset / Migration → ТП дороже на 10–15% · I2 More-for-More · I3 Бустер при исчерпании + роуминг-промо · I4 Loyalty (кэшбэк, Daily Drops, Weekly Wins, Monthly Grand)", None),
        ("", None),
        ("Как пользоваться", BOLD),
        ("1. Жёлтые ячейки — гипотезы: заменить фактом из запроса данных (v2/docs/12_data_request_for_analysts.md).", None),
        ("2. Синие — параметры (меняйте по итогам пилотов и MVP). Чёрные — формулы. Зелёные — ссылки на другие листы.", None),
        ("3. Monthly_I и Loyalty_Monthly — помесячные движки на 36 мес. (M1 = ноябрь 2026).", None),
        ("4. Prizes и Daily_Drops — каталог призов в сумах, фонды Weekly/Monthly, таблица вероятностей сундука.", None),
        ("5. Thresholds — сравнение вариантов порога участия (O1–O5). MVP_Budget — бюджет пилота программы.", None),
        ("6. MC_Results — вероятности из Монте-Карло (python v2/model/model_v2.py), это значения, не формулы.", None),
        ("", None),
        ("Единицы: сумы без НДС 12% (цены призов и абонплаты — с НДС, где указано). ARPU — по когорте архивных абонентов T0.", None),
        ("Excel повторяет детерминированный прогон Python на модальных значениях.", None),
    ]
    for i, (t, f) in enumerate(lines, 1):
        ws.cell(row=i, column=1, value=t).font = f or BLACK
    ws.column_dimensions["A"].width = 160


# ============================================================================ Inputs
INP = {}


def sheet_inputs(wb):
    ws = wb.create_sheet("Inputs")
    ws["A1"] = "Базовые допущения (архивная когорта)"
    ws["A1"].font = TITLE
    for j, h in enumerate(["Параметр", "", "Значение", "Ед.", "Источник / комментарий"], 1):
        ws.cell(row=3, column=j, value=h)
    hdr(ws, 3, 5)
    b = A.BASE
    rows = [
        ("total", "Абонентская база Ucell", b["total_subs"], "абон.", "Ucell: >10,9 млн на начало 2026", False),
        ("ashare", "Доля абонентов на архивных ТП", b["archive_share"], "%", "HYPOTHESIS — запрос D1", True),
        ("aactive", "Доля активных среди архивных", b["archive_active_ratio"], "%", "HYPOTHESIS — запрос D1", True),
        ("N", "Активная архивная когорта (N)", "=C4*C5*C6", "абон.", "Формула", False),
        ("ARPU0", "ARPU архивной когорты (net)", b["archive_arpu"], "сум/мес", "HYPOTHESIS — запрос D2", True),
        ("c", "Месячный отток архивной когорты", b["monthly_churn"], "%", "HYPOTHESIS — запрос D4", True),
        ("VAT", "НДС", A.VAT, "%", "НК РУз: 12%", False),
        ("FX", "Курс, сум за 1 USD", A.FX_UZS_PER_USD, "сум", "ЦБ РУз, 2026", False),
        ("MKT", "Рынок мобильной связи, абонентов", b["market_total_subs"], "абон.", "HYPOTHESIS", True),
        ("R0m", "Выручка когорты в месяц", "=C7*C8", "сум/мес", "Формула", False),
        ("h", "Перекрытие эффектов I1–I3", mode(A.PORTFOLIO_OVERLAP_HAIRCUT), "%", "Одни и те же абоненты в нескольких инициативах", True),
    ]
    for i, (k, name, v, unit, src, key) in enumerate(rows, 4):
        ws.cell(row=i, column=1, value=name)
        is_f = isinstance(v, str) and v.startswith("=")
        put(ws, i, 3, v, BLACK if is_f else BLUE, PCT if unit == "%" else NUM, YELLOW if key else None)
        ws.cell(row=i, column=4, value=unit)
        ws.cell(row=i, column=5, value=src)
        INP[k] = f"Inputs!$C${i}"
    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["C"].width = 16
    ws.column_dimensions["E"].width = 60


# ============================================================================ Initiatives + Monthly_I
IPARAMS = [
    ("name", "Название", None), ("short", "Коротко", None), ("type", "Тип механики", None),
    ("complexity", "Сложность биллинг/интеграции (1–5)", NUM),
    ("segment_arpu", "ARPU целевого сегмента, сум", NUM),
    ("eligible_share", "Доля целевой группы от когорты", PCT),
    ("price_up_pct", "Новый ТП дороже на, % (I1)", PCT),
    ("base_fee_gross", "Текущая абонплата мигрируемых линеек, с НДС (I1)", NUM),
    ("price_up_gross", "Повышение абонплаты, сум с НДС (I2)", NUM),
    ("upsell_share", "Доля выбравших ТП крупнее", PCT),
    ("upsell_delta", "Доп. ARPU у выбравших крупнее, сум", NUM),
    ("downgrade", "Доля даунгрейда", PCT), ("downgrade_delta", "Δ выручки при даунгрейде, сум", NUM),
    ("churn_incr", "Доп. отток", PCT), ("data_cost_m", "Себестоимость доп. трафика, сум/мес", NUM),
    ("opex_savings_m", "Экономия OPEX, сум/мес", NUM),
    ("reach", "Охват", PCT), ("takeup", "Конверсия", PCT), ("cannibal", "Каннибализация", PCT),
    ("delta", "ΔARPU на адоптера, сум", NUM), ("churn_red", "Снижение оттока адоптеров", PCT),
    ("cost_per_contact", "Стоимость контакта, сум", NUM),
    ("cost_per_adopter_once", "Разовая стоимость на адоптера, сум", NUM),
    ("var_cost_m", "Перем. затраты на адоптера, сум/мес", NUM),
    ("one_time", "Разовые затраты, сум", NUM), ("opex_m", "OPEX, сум/мес", NUM),
    ("margin", "Маржинальность", PCT),
    ("L1", "Месяц запуска / волна 1", NUM), ("frac1", "Доля волны 1 (reprice)", PCT),
    ("L2", "Месяц волны 2 (reprice)", NUM), ("ramp", "Набор, мес. (adoption)", NUM),
    ("exec_prob", "Вероятность запуска", PCT),
]
IDER = [
    ("elig", "Целевая группа, абон.", NUM), ("contacted", "Охвачено, абон.", NUM),
    ("a_final", "Адоптеры при полном наборе", NUM), ("up_net", "Повышение абонплаты net, сум", NUM),
    ("stay", "Остаются и платят больше", NUM), ("down", "Даунгрейд, абон.", NUM),
    ("lost_total", "Доп. отток всего, абон.", NUM), ("build", "Месяцев подготовки", NUM),
    ("eval_m", "Месяц оценки (12-й после старта)", NUM),
]
ISUM = [
    ("rev1", "Выручка Y1", NUM), ("rev2", "Выручка Y2", NUM), ("rev3", "Выручка Y3", NUM),
    ("runrate", "Run-rate/год на M24", NUM), ("cost24", "Затраты 24 мес.", NUM),
    ("contrib24", "Вклад 24 мес.", NUM), ("contrib36", "Вклад 36 мес.", NUM),
    ("roi24", "ROI 24 мес.", MULT), ("upl_eval", "ΔARPU когорты на месяц оценки", PCT2),
    ("upl24", "ΔARPU когорты на M24", PCT2), ("dsubs24", "Δ абонентов на M24", NUM),
    ("exp24", "Ожидаемый вклад 24м × P(запуска)", NUM),
]
IBLOCK = ["f", "target", "new", "R", "pool", "w", "newlost", "L", "rev", "cost", "contrib", "dsubs"]
IBL = {"f": "Доля набора", "target": "Адоптеры (цель)", "new": "Новые адоптеры", "R": "Удержанные",
       "pool": "Пул адоптеров", "w": "Вес волн × затухание", "newlost": "Новый доп. отток", "L": "Потерянные",
       "rev": "Инкр. выручка", "cost": "Затраты", "contrib": "Вклад", "dsubs": "Δ абонентов"}


def ival(p, key):
    if key == "L1":
        return p["waves"][0][0] if p["type"] == "reprice" else p["launch"]
    if key == "frac1":
        return p["waves"][0][1] if p["type"] == "reprice" else 1
    if key == "L2":
        return p["waves"][1][0] if p["type"] == "reprice" and len(p["waves"]) > 1 else 99
    if key in ("name", "short", "type"):
        return p.get(key, "")
    if key == "ramp":
        return p.get("ramp", 1)
    return mode(p.get(key, 0))


def sheets_initiatives(wb):
    keys = list(A.INITIATIVES)
    wi = wb.create_sheet("Initiatives")
    wi["A1"] = "Инициативы I1–I3: параметры (синие) → производные → итоги (из Monthly_I)"
    wi["A1"].font = TITLE
    wi.cell(row=3, column=1, value="Параметр")
    wi.cell(row=3, column=2, value="Ключ")
    for j, k in enumerate(keys, 3):
        wi.cell(row=3, column=j, value=k)
    hdr(wi, 3, 2 + len(keys))
    row = {}
    r = 4
    for key, label, fmt in IPARAMS:
        row[key] = r
        wi.cell(row=r, column=1, value=label)
        wi.cell(row=r, column=2, value=key)
        for j, k in enumerate(keys, 3):
            p = A.INITIATIVES[k]
            c = put(wi, r, j, ival(p, key), BLUE, fmt)
            if isinstance(p.get(key), (tuple, list)):
                c.comment = Comment(f"Монте-Карло: {p[key][0]:,} … {p[key][2]:,}", "model")
            if key == "name":
                c.alignment = Alignment(wrap_text=True, vertical="top")
        r += 1
    r += 1
    section(wi, r, "ПРОИЗВОДНЫЕ", 2 + len(keys))
    r += 1
    for key, label, fmt in IDER:
        row[key] = r
        wi.cell(row=r, column=1, value=label)
        wi.cell(row=r, column=2, value=key)
        r += 1
    r += 1
    section(wi, r, "ИТОГИ", 2 + len(keys))
    r += 1
    for key, label, fmt in ISUM:
        row[key] = r
        wi.cell(row=r, column=1, value=label)
        wi.cell(row=r, column=2, value=key)
        r += 1

    wm = wb.create_sheet("Monthly_I")
    wm["A1"] = "Помесячный движок I1–I3 (36 мес., сум без НДС)"
    wm["A1"].font = TITLE
    wm.cell(row=3, column=1, value="Месяц")
    for m in range(1, H + 1):
        wm.cell(row=3, column=M0 + m - 1, value=m)
    hdr(wm, 3, M0 + H - 1)
    wm.cell(row=4, column=1, value="S0: абоненты когорты без инициатив")
    wm.cell(row=5, column=1, value="R0: выручка когорты без инициатив")
    for m in range(1, H + 1):
        X = mc(m)
        put(wm, 4, M0 + m - 1, f"={INP['N']}*(1-{INP['c']})^{X}$3", fmt=NUM)
        put(wm, 5, M0 + m - 1, f"={X}4*{INP['ARPU0']}", fmt=NUM)
    brow = {}
    r = 7
    for k in keys:
        section(wm, r, f"{k} — {A.INITIATIVES[k]['name']}", M0 + H - 1)
        r += 1
        for it in IBLOCK:
            brow[(k, it)] = r
            wm.cell(row=r, column=1, value=IBL[it])
            wm.cell(row=r, column=2, value=it)
            r += 1
        r += 1
    c = INP["c"]
    for j, k in enumerate(keys, 3):
        J = CL(j)

        def P(key):
            return f"Initiatives!${J}${row[key]}"

        for m in range(1, H + 1):
            X, W = mc(m), CL(M0 + m - 2)
            mm = f"{X}$3"

            def B(it, col=X):
                return f"{col}{brow[(k, it)]}"

            ad = f'{P("type")}="adoption"'
            f = {
                "f": f"=IF({ad},IF({mm}<{P('L1')},0,MIN(({mm}-{P('L1')}+1)/{P('ramp')},1)),0)",
                "target": f"={P('a_final')}*{B('f')}",
                "new": f"=MAX({B('target')}-N({B('target', W)}),0)",
                "R": f"=IF({ad},N({B('R', W)})*(1-{c})+N({B('pool', W)})*{c}*{P('churn_red')},0)",
                "pool": f"=N({B('pool', W)})*(1-{c}*(1-{P('churn_red')}))+{B('new')}",
                "w": (f"=IF({ad},0,IF({mm}>={P('L1')},{P('frac1')}*(1-{c})^({mm}-{P('L1')}),0)"
                      f"+IF({mm}>={P('L2')},(1-{P('frac1')})*(1-{c})^({mm}-{P('L2')}),0))"),
                "newlost": (f"=IF({ad},IF(AND({mm}>={P('L1')},{mm}<{P('L1')}+3),{P('lost_total')}/3,0),"
                            f"IF(AND({mm}>={P('L1')},{mm}<{P('L1')}+2),{P('frac1')}*{P('lost_total')}/2,0)"
                            f"+IF(AND({mm}>={P('L2')},{mm}<{P('L2')}+2),(1-{P('frac1')})*{P('lost_total')}/2,0))"),
                "L": f"=N({B('L', W)})*(1-{c})+{B('newlost')}",
                "rev": (f"=IF({ad},{B('pool')}*{P('delta')}+{B('R')}*{P('segment_arpu')}-{B('L')}*{P('segment_arpu')},"
                        f"({P('stay')}*{P('up_net')}+{P('elig')}*{P('upsell_share')}*{P('upsell_delta')}"
                        f"+{P('down')}*{P('downgrade_delta')})*{B('w')}-{B('L')}*{P('segment_arpu')})"),
                "cost": (f"=IF({mm}<={P('build')},{P('one_time')}/{P('build')},0)+IF({ad},"
                         f"{B('new')}*{P('cost_per_adopter_once')}+{B('pool')}*{P('var_cost_m')}"
                         f"+IF({mm}={P('L1')},{P('contacted')}*{P('cost_per_contact')},0)+IF({mm}>={P('L1')},{P('opex_m')},0),"
                         f"{P('stay')}*{P('data_cost_m')}*{B('w')}+IF({mm}={P('L1')},{P('frac1')}*{P('elig')}*{P('cost_per_contact')},0)"
                         f"+IF({mm}={P('L2')},(1-{P('frac1')})*{P('elig')}*{P('cost_per_contact')},0)"
                         f"+IF({mm}>={P('L1')},{P('opex_m')}-{P('opex_savings_m')}*({P('frac1')}+IF({mm}>={P('L2')},1-{P('frac1')},0)),0))"),
                "contrib": f"={B('rev')}*{P('margin')}-{B('cost')}",
                "dsubs": f"={B('R')}-{B('L')}",
            }
            for it in IBLOCK:
                put(wm, brow[(k, it)], M0 + m - 1, f[it], fmt=PCT if it == "f" else NUM)

        der = {
            "elig": f"={INP['N']}*{J}{row['eligible_share']}",
            "contacted": f'=IF({J}{row["type"]}="adoption",{J}{row["elig"]}*{J}{row["reach"]},{J}{row["elig"]})',
            "a_final": f'=IF({J}{row["type"]}="adoption",{J}{row["contacted"]}*{J}{row["takeup"]}*(1-{J}{row["cannibal"]}),0)',
            "up_net": (f"=IF({J}{row['price_up_pct']}>0,{J}{row['base_fee_gross']}*{J}{row['price_up_pct']},"
                       f"{J}{row['price_up_gross']})/(1+{INP['VAT']})"),
            "stay": f'=IF({J}{row["type"]}="reprice",{J}{row["elig"]}*(1-{J}{row["downgrade"]}-{J}{row["churn_incr"]}),0)',
            "down": f'=IF({J}{row["type"]}="reprice",{J}{row["elig"]}*{J}{row["downgrade"]},0)',
            "lost_total": f"={J}{row['contacted']}*{J}{row['churn_incr']}",
            "build": f"=MAX({J}{row['L1']}-1,1)",
            "eval_m": f"=MIN({J}{row['L1']}+11,{H})",
        }
        for key, label, fmt in IDER:
            put(wi, row[key], j, der[key], BLACK, fmt)

        def rng(it, a, b):
            rr = brow[(k, it)]
            return f"Monthly_I!{mc(a)}{rr}:{mc(b)}{rr}"

        def at(it, expr):
            rr = brow[(k, it)]
            return f"INDEX(Monthly_I!${mc(1)}${rr}:${mc(H)}${rr},1,{expr})"

        e = f"{J}{row['eval_m']}"
        summ = {
            "rev1": f"=SUM({rng('rev', 1, 12)})", "rev2": f"=SUM({rng('rev', 13, 24)})",
            "rev3": f"=SUM({rng('rev', 25, 36)})", "runrate": f"=Monthly_I!{mc(24)}{brow[(k, 'rev')]}*12",
            "cost24": f"=SUM({rng('cost', 1, 24)})", "contrib24": f"=SUM({rng('contrib', 1, 24)})",
            "contrib36": f"=SUM({rng('contrib', 1, 36)})",
            "roi24": f"=IF({J}{row['cost24']}>0,{J}{row['contrib24']}/{J}{row['cost24']},0)",
            "upl_eval": (f"=(INDEX(Monthly_I!${mc(1)}$5:${mc(H)}$5,1,{e})+{at('rev', e)})"
                         f"/(INDEX(Monthly_I!${mc(1)}$4:${mc(H)}$4,1,{e})+{at('dsubs', e)})/{INP['ARPU0']}-1"),
            "upl24": (f"=(Monthly_I!{mc(24)}$5+Monthly_I!{mc(24)}{brow[(k, 'rev')]})"
                      f"/(Monthly_I!{mc(24)}$4+Monthly_I!{mc(24)}{brow[(k, 'dsubs')]})/{INP['ARPU0']}-1"),
            "dsubs24": f"=Monthly_I!{mc(24)}{brow[(k, 'dsubs')]}",
            "exp24": f"={J}{row['contrib24']}*{J}{row['exec_prob']}",
        }
        for key, label, fmt in ISUM:
            put(wi, row[key], j, summ[key], GREEN if key not in ("roi24", "exp24") else BLACK, fmt)
    wi.column_dimensions["A"].width = 52
    wi.column_dimensions["B"].width = 18
    for j in range(3, 3 + len(keys)):
        wi.column_dimensions[CL(j)].width = 22
    wi.row_dimensions[row["name"]].height = 70
    wi.freeze_panes = "C4"
    wm.column_dimensions["A"].width = 40
    for m in range(1, H + 1):
        wm.column_dimensions[mc(m)].width = 15
    wm.freeze_panes = "C4"
    return keys, row, brow


# ============================================================================ Loyalty inputs, prizes, drops
LI = {}


def sheet_loyalty_inputs(wb):
    ws = wb.create_sheet("Loyalty_Inputs")
    ws["A1"] = "Loyalty program: параметры"
    ws["A1"].font = TITLE
    for j, h in enumerate(["Параметр", "Ключ", "Значение", "Комментарий"], 1):
        ws.cell(row=3, column=j, value=h)
    hdr(ws, 3, 4)
    L = A.LOYALTY
    items = [
        ("app_mau_share", "MAU приложения / база", PCT, True, "HYPOTHESIS — запрос L1"),
        ("enroll_of_mau", "Вступили в программу из MAU", PCT, True, "HYPOTHESIS; проверяется MVP"),
        ("archive_share_members", "Доля архивных среди участников", PCT, True, "HYPOTHESIS — запрос L1"),
        ("mvp_participants", "Участников MVP", NUM, False, "Стратифицированная выборка"),
        ("mvp_start", "Месяц старта MVP (M5 = март 2027)", NUM, False, ""),
        ("mvp_months", "Длительность MVP, мес.", NUM, False, ""),
        ("scale_start", "Месяц старта масштаба (M11 = сентябрь 2027)", NUM, False, ""),
        ("scale_ramp", "Набор участников при масштабе, мес.", NUM, False, ""),
        ("ontime_share", "Доля оплат вовремя (кэшбэк только им)", PCT, True, "запрос L3"),
        ("breakage", "Breakage: монеты не потрачены (12 мес.)", PCT, False, "Бенчмарки 25–40%"),
        ("redeem_cost_ratio", "Себестоимость погашения монет", PCT, False, "50% услуги×0,15 + 30% счёт×1 + 20% мерч×0,7"),
        ("drops_dau_share", "Ежедневно открывают сундук (доля участников)", PCT, True, "проверяется MVP"),
        ("mvp_prize_share", "Призовые фонды MVP, доля от масштаба", PCT, False, ""),
        ("mvp_build", "Разработка MVP, сум", NUM, False, "web-view в приложении"),
        ("scale_build", "Разработка масштаба, сум", NUM, False, ""),
        ("mvp_opex_m", "OPEX MVP, сум/мес", NUM, False, ""),
        ("opex_m", "OPEX программы, сум/мес", NUM, False, "команда, лицензии, фулфилмент, КЦ"),
        ("marketing_m", "Маркетинг, сум/мес", NUM, False, ""),
        ("launch_marketing", "Маркетинг запуска масштаба, сум", NUM, False, ""),
        ("mvp_marketing", "Маркетинг MVP, сум", NUM, False, ""),
        ("member_churn_base", "Месячный отток участников (база)", PCT2, True, "запрос L4"),
        ("churn_reduction", "Снижение оттока участников (отн.)", PCT, False, "T-Mobile Tuesdays −27% (с самоотбором)"),
        ("ontime_rev_uplift", "Рост выручки от оплат вовремя", PCT, False, "H2 MVP"),
        ("upgrade_rate_m", "Доп. переходы на ТП ≥ 60 тыс., в мес.", PCT2, False, "H1 MVP — главный драйвер"),
        ("upgrade_delta", "ΔARPU при переходе, сум", NUM, False, ""),
        ("pass_adoption", "Club+ Pass: доля участников < 60 тыс.", PCT, False, "H3 MVP"),
        ("pass_price_gross", "Club+ Pass: цена с НДС, сум/мес", NUM, False, ""),
        ("pass_value_cost", "Себестоимость ценности Pass", PCT, False, "ГБ, ×1,5 кэшбэк"),
        ("engagement_upsell", "Допродажи в приложении (при DAU 30%)", PCT2, False, ""),
        ("partner_cash_m", "Платные размещения партнёров, сум/мес (с M+6 масштаба)", NUM, False, ""),
        ("margin", "Маржинальность доп. выручки", PCT, False, ""),
        ("archive_member_arpu", "ARPU архивного участника, сум", NUM, True, "запрос L1"),
        ("archive_churn_base", "Отток архивных участников", PCT2, True, ""),
        ("archive_upgrade_rate_m", "Архивные: переход на актуальный ТП, в мес.", PCT2, False, ""),
        ("archive_upgrade_delta", "Архивные: ΔARPU при переходе, сум", NUM, False, ""),
    ]
    r = 4
    for k, label, fmt, key, cm in items:
        ws.cell(row=r, column=1, value=label)
        ws.cell(row=r, column=2, value=k)
        v = L[k]
        c = put(ws, r, 3, mode(v), BLUE, fmt, YELLOW if key else None)
        if isinstance(v, (tuple, list)):
            c.comment = Comment(f"Монте-Карло: {v[0]:,} … {v[2]:,}", "model")
        ws.cell(row=r, column=4, value=cm)
        LI[k] = f"Loyalty_Inputs!$C${r}"
        r += 1
    r += 1
    ws.cell(row=r, column=1, value="Распределение участников по абонплате (HYPOTHESIS — запрос L2)").font = BOLD
    r += 1
    heads = ["Полоса", "Ключ", "Доля участников", "Доля архивных участников", "Средняя АП с НДС", "Кэшбэк", "Шансы/нед. (O5)"]
    for j, h in enumerate(heads, 1):
        ws.cell(row=r, column=j, value=h)
    hdr(ws, r, len(heads))
    r += 1
    b0 = r
    for k, v in A.BANDS.items():
        put(ws, r, 1, v["name"])
        put(ws, r, 2, k)
        put(ws, r, 3, v["member_share"], BLUE, PCT, YELLOW)
        put(ws, r, 4, v["archive_share"], BLUE, PCT, YELLOW)
        put(ws, r, 5, v["fee"], BLUE, NUM)
        put(ws, r, 6, v["cb"], BLUE, PCT)
        put(ws, r, 7, v["entries"], BLUE, NUM)
        r += 1
    b1 = r - 1
    r += 1
    der = [
        ("arpu_net", "ARPU участника net", f"=SUMPRODUCT(C{b0}:C{b1},E{b0}:E{b1})/(1+{INP['VAT']})", NUM),
        ("earn", "Начисляется монет на участника/мес", f"=SUMPRODUCT(C{b0}:C{b1},E{b0}:E{b1},F{b0}:F{b1})", NUM),
        ("cb_avg", "Средняя ставка кэшбэка", f"=C{r+1}/SUMPRODUCT(C{b0}:C{b1},E{b0}:E{b1})", PCT2),
        ("below60", "Доля участников с АП < 60 тыс.", f"=C{b0}+C{b0+1}", PCT),
        ("a_earn", "Монет на архивного участника/мес", f"=SUMPRODUCT(D{b0}:D{b1},E{b0}:E{b1},F{b0}:F{b1})", NUM),
        ("a_below60", "Доля архивных участников < 60 тыс.", f"=D{b0}+D{b0+1}", PCT),
        ("full", "Участников при полном масштабе", f"={INP['total']}*{LI['app_mau_share']}*{LI['enroll_of_mau']}", NUM),
        ("pass_net", "Club+ Pass net, сум", f"={LI['pass_price_gross']}/(1+{INP['VAT']})", NUM),
        ("check", "Проверка долей", f'=IF(AND(ABS(SUM(C{b0}:C{b1})-1)<0.0001,ABS(SUM(D{b0}:D{b1})-1)<0.0001),"OK","ОШИБКА")', None),
    ]
    for k, label, f, fmt in der:
        ws.cell(row=r, column=1, value=label)
        ws.cell(row=r, column=2, value=k)
        put(ws, r, 3, f, BLACK, fmt)
        LI[k] = f"Loyalty_Inputs!$C${r}"
        r += 1
    ws.column_dimensions["A"].width = 52
    ws.column_dimensions["B"].width = 24
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 46
    for col in "EFG":
        ws.column_dimensions[col].width = 16


def sheet_prizes(wb):
    ws = wb.create_sheet("Prizes")
    ws["A1"] = "Каталог призов (цены — ориентиры 10/2026, уточнить у закупок) и призовые фонды"
    ws["A1"].font = TITLE
    for j, h in enumerate(["Ключ", "Приз", "Цена, сум с НДС", "Источник"], 1):
        ws.cell(row=3, column=j, value=h)
    hdr(ws, 3, 4)
    r = 4
    p0 = r
    for k, v in A.PRIZES.items():
        put(ws, r, 1, k)
        put(ws, r, 2, v["name"])
        put(ws, r, 3, v["price"], BLUE, NUM, YELLOW)
        put(ws, r, 4, v["src"])
        r += 1
    p1 = r - 1
    cat = f"$A${p0}:$A${p1}"
    pr = f"$C${p0}:$C${p1}"

    def table(r, title, rows, key):
        ws.cell(row=r, column=1, value=title).font = BOLD
        r += 1
        for j, h in enumerate(["Ключ", "Приз", "Кол-во в неделю", "Софинансирует партнёр", "Затраты/нед., сум"], 1):
            ws.cell(row=r, column=j, value=h)
        hdr(ws, r, 5)
        r += 1
        t0 = r
        for k, q, co in rows:
            put(ws, r, 1, k)
            put(ws, r, 2, f"=INDEX($B${p0}:$B${p1},MATCH(A{r},{cat},0))", GREEN)
            put(ws, r, 3, q, BLUE, NUM)
            put(ws, r, 4, co, BLUE, PCT)
            put(ws, r, 5, f"=INDEX({pr},MATCH(A{r},{cat},0))*C{r}*(1-D{r})", BLACK, NUM)
            r += 1
        t1 = r - 1
        ws.cell(row=r, column=2, value="Итого в неделю").font = BOLD
        put(ws, r, 3, f"=SUM(C{t0}:C{t1})", BOLD, NUM)
        put(ws, r, 5, f"=SUM(E{t0}:E{t1})", BOLD, NUM)
        LI[key] = f"Prizes!$E${r}"
        LI[key + "_items"] = f"Prizes!$C${r}"
        ws.cell(row=r + 1, column=2, value="Итого в месяц (×52/12)")
        put(ws, r + 1, 5, f"=E{r}*52/12", BOLD, NUM)
        return r + 3

    r = table(r + 1, "Weekly Wins — еженедельный фонд (масштаб)", A.WEEKLY_WINS, "ww_week")
    r = table(r, "Daily Drops — еженедельный инвентарь (масштаб)", A.DROPS_INVENTORY, "inv_week")
    ws.cell(row=r, column=1, value="Monthly Grand — календарь на 12 месяцев").font = BOLD
    r += 1
    for j, h in enumerate(["Месяц", "Приз", "Стоимость, сум"], 1):
        ws.cell(row=r, column=j, value=h)
    hdr(ws, r, 3)
    r += 1
    g0 = r
    months = ["Сен", "Окт", "Ноя", "Дек", "Янв", "Фев", "Мар", "Апр", "Май", "Июн", "Июл", "Авг"]
    for i, g in enumerate(A.MONTHLY_GRAND):
        put(ws, r, 1, months[i])
        if g == "iphone17_x5":
            put(ws, r, 2, "5 × iPhone 17")
            put(ws, r, 3, f"=5*INDEX({pr},MATCH(\"iphone17\",{cat},0))", BLACK, NUM)
        elif g == "free_year_x10":
            put(ws, r, 2, "10 × год связи бесплатно")
            put(ws, r, 3, f"=10*INDEX({pr},MATCH(\"free_year\",{cat},0))", BLACK, NUM)
        else:
            put(ws, r, 2, A.PRIZES[g]["name"])
            put(ws, r, 3, f"=INDEX({pr},MATCH(\"{g}\",{cat},0))", BLACK, NUM)
        r += 1
    g1 = r - 1
    ws.cell(row=r, column=2, value="Итого за год").font = BOLD
    put(ws, r, 3, f"=SUM(C{g0}:C{g1})", BOLD, NUM)
    ws.cell(row=r + 1, column=2, value="Накладные (логистика, нотариус/комиссия, резерв налогов)")
    put(ws, r + 1, 3, A.GRAND_OVERHEAD, BLUE, PCT)
    ws.cell(row=r + 2, column=2, value="В среднем в месяц с накладными").font = BOLD
    put(ws, r + 2, 3, f"=C{r}/12*(1+C{r+1})", BOLD, NUM)
    LI["grand_month"] = f"Prizes!$C${r+2}"
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 52
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 30
    ws.column_dimensions["E"].width = 20


def sheet_drops(wb):
    ws = wb.create_sheet("Daily_Drops")
    ws["A1"] = "Daily Drops: что выпадает из сундука (цифровые призы) и недельный бюджет"
    ws["A1"].font = TITLE
    for j, h in enumerate(["Приз", "Вероятность", "Себестоимость, сум", "Ожидаемо на открытие, сум", "Комментарий"], 1):
        ws.cell(row=3, column=j, value=h)
    hdr(ws, 3, 5)
    r = 4
    for name, prob, cost, cm in A.DROP_TABLE:
        put(ws, r, 1, name)
        put(ws, r, 2, prob, BLUE, PCT)
        put(ws, r, 3, cost, BLUE, NUM1)
        put(ws, r, 4, f"=B{r}*C{r}", BLACK, NUM1)
        put(ws, r, 5, cm)
        r += 1
    e = r - 1
    ws.cell(row=r, column=1, value="Итого").font = BOLD
    put(ws, r, 2, f"=SUM(B4:B{e})", BOLD, PCT)
    put(ws, r, 4, f"=SUM(D4:D{e})", BOLD, NUM1)
    LI["ddc"] = f"Daily_Drops!$D${r}"
    ws.cell(row=r + 1, column=1, value="Проверка: вероятности = 100%")
    put(ws, r + 1, 2, f'=IF(ABS(B{r}-1)<0.0001,"OK","ОШИБКА")')
    r += 3
    ws.cell(row=r, column=1, value="Недельный бюджет Daily Drops при масштабе").font = BOLD
    r += 1
    rows = [
        ("Участников", f"={LI['full']}", NUM),
        ("Открытий в неделю", f"=C{r}*{LI['drops_dau_share']}*7", NUM),
        ("Цифровые призы, сум/нед.", f"=C{r+1}*{LI['ddc']}", NUM),
        ("Инвентарь, сум/нед. (лист Prizes)", f"={LI['inv_week']}", NUM),
        ("Итого Daily Drops, сум/нед.", f"=C{r+2}+C{r+3}", NUM),
        ("Стоимость на одно открытие, сум", f"=C{r+4}/C{r+1}", NUM1),
        ("Правило безубыточности: ≤ доп. маржа от допродаж на одно открытие, сум",
         f"={LI['arpu_net']}*{LI['engagement_upsell']}*{LI['margin']}/({LI['drops_dau_share']}*30)*({LI['drops_dau_share']}/0.3)", NUM1),
    ]
    for label, f, fmt in rows:
        ws.cell(row=r, column=1, value=label)
        put(ws, r, 3, f, BLACK, fmt)
        r += 1
    ws.column_dimensions["A"].width = 64
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 20
    ws.column_dimensions["D"].width = 24
    ws.column_dimensions["E"].width = 50


# ============================================================================ Loyalty monthly
LBLOCK = [
    ("mem", "Участники", NUM), ("pf", "Коэф. призовых фондов", PCT),
    ("c_build", "Разработка", NUM), ("c_opex", "OPEX", NUM), ("c_mkt", "Маркетинг", NUM),
    ("c_cb", "Кэшбэк (себестоимость погашений)", NUM), ("c_dd", "Daily Drops — цифровые", NUM),
    ("c_inv", "Daily Drops — инвентарь", NUM), ("c_ww", "Weekly Wins", NUM), ("c_gr", "Monthly Grand", NUM),
    ("cost", "ЗАТРАТЫ ВСЕГО", NUM),
    ("R", "Удержанные участники", NUM), ("b_ret", "Выручка: удержание", NUM), ("b_ont", "Выручка: оплаты вовремя", NUM),
    ("U", "Накопленные переходы на ТП ≥ 60 тыс.", NUM), ("b_upg", "Выручка: переходы", NUM),
    ("b_pass", "Club+ Pass (net)", NUM), ("b_eng", "Выручка: допродажи", NUM), ("b_par", "Партнёры (платные размещения)", NUM),
    ("contrib", "ВКЛАД", NUM), ("cum", "Вклад накопленный", NUM),
    ("am", "Архивные участники", NUM), ("Ra", "Архивные: удержанные", NUM), ("Ua", "Архивные: перешли на актуальный ТП", NUM),
    ("contra", "Архивные: погашение счёта монетами", NUM), ("a_rev", "Архивные: доп. выручка", NUM), ("a_ds", "Архивные: Δ абонентов", NUM),
]


def sheet_loyalty_monthly(wb):
    ws = wb.create_sheet("Loyalty_Monthly")
    ws["A1"] = "Loyalty program: помесячный движок (36 мес.; MVP M5–M7, масштаб с M11)"
    ws["A1"].font = TITLE
    ws.cell(row=3, column=1, value="Месяц")
    for m in range(1, H + 1):
        ws.cell(row=3, column=M0 + m - 1, value=m)
    hdr(ws, 3, M0 + H - 1)
    lr = {}
    r = 4
    for k, label, fmt in LBLOCK:
        lr[k] = r
        ws.cell(row=r, column=1, value=label).font = BOLD if k in ("cost", "contrib") else BLACK
        ws.cell(row=r, column=2, value=k)
        r += 1
    L = LI
    for m in range(1, H + 1):
        X, W = mc(m), CL(M0 + m - 2)
        mm = f"{X}$3"

        def B(k, col=X):
            return f"{col}{lr[k]}"

        ms, mmn, ss, sr = L["mvp_start"], L["mvp_months"], L["scale_start"], L["scale_ramp"]
        f = {
            "mem": f"=IF({mm}<{ms},0,IF({mm}<{ss},{L['mvp_participants']},MAX({L['mvp_participants']},{L['full']}*MIN(({mm}-{ss}+1)/{sr},1))))",
            "pf": f"=IF({B('mem')}>0,MAX({L['mvp_prize_share']},{B('mem')}/{L['full']}),0)",
            "c_build": (f"=IF(AND({mm}>=2,{mm}<{ms}),{L['mvp_build']}/({ms}-2),0)"
                        f"+IF(AND({mm}>={ms}+{mmn},{mm}<{ss}),{L['scale_build']}/({ss}-{ms}-{mmn}),0)"),
            "c_opex": f"=IF(AND({mm}>={ms},{mm}<{ss}),{L['mvp_opex_m']},0)+IF({mm}>={ss},{L['opex_m']},0)",
            "c_mkt": (f"=IF({mm}>={ss},{L['marketing_m']},0)+IF(OR({mm}={ss},{mm}={ss}+1),{L['launch_marketing']}/2,0)"
                      f"+IF({mm}={ms},{L['mvp_marketing']},0)"),
            "c_cb": f"={B('mem')}*{L['ontime_share']}*{L['earn']}*(1-{L['breakage']})*{L['redeem_cost_ratio']}",
            "c_dd": f"={B('mem')}*{L['drops_dau_share']}*30*{L['ddc']}",
            "c_inv": f"={L['inv_week']}*52/12*{B('pf')}",
            "c_ww": f"={L['ww_week']}*52/12*{B('pf')}",
            "c_gr": f"={L['grand_month']}*{B('pf')}",
            "cost": f"=SUM({X}{lr['c_build']}:{X}{lr['c_gr']})",
            "R": f"=N({B('R', W)})*(1-{L['member_churn_base']})+N({B('mem', W)})*{L['member_churn_base']}*{L['churn_reduction']}",
            "b_ret": f"={B('R')}*{L['arpu_net']}",
            "b_ont": f"={B('mem')}*{L['arpu_net']}*{L['ontime_rev_uplift']}",
            "U": f"=N({B('U', W)})*(1-{L['member_churn_base']})+{B('mem')}*{L['below60']}*{L['upgrade_rate_m']}",
            "b_upg": f"={B('U')}*{L['upgrade_delta']}",
            "b_pass": f"={B('mem')}*{L['below60']}*{L['pass_adoption']}*{L['pass_net']}",
            "b_eng": f"={B('mem')}*{L['arpu_net']}*{L['engagement_upsell']}*{L['drops_dau_share']}/0.3",
            "b_par": f"=IF({mm}>={ss}+6,{L['partner_cash_m']},0)",
            "contrib": (f"=({B('b_ret')}+{B('b_ont')}+{B('b_upg')}+{B('b_eng')})*{L['margin']}"
                        f"+{B('b_pass')}*(1-{L['pass_value_cost']})+{B('b_par')}-{B('cost')}"),
            "cum": f"=N({B('cum', W)})+{B('contrib')}",
            "am": f"={B('mem')}*{L['archive_share_members']}",
            "Ra": f"=N({B('Ra', W)})*(1-{L['archive_churn_base']})+N({B('am', W)})*{L['archive_churn_base']}*{L['churn_reduction']}",
            "Ua": f"=N({B('Ua', W)})*(1-{L['archive_churn_base']})+{B('am')}*{L['archive_upgrade_rate_m']}",
            "contra": f"={B('am')}*{L['ontime_share']}*{L['a_earn']}*(1-{L['breakage']})*0.3",
            "a_rev": (f"={B('Ra')}*{L['archive_member_arpu']}+{B('Ua')}*{L['archive_upgrade_delta']}"
                      f"+{B('am')}*{L['archive_member_arpu']}*{L['ontime_rev_uplift']}"
                      f"+{B('am')}*{L['a_below60']}*{L['pass_adoption']}*{L['pass_net']}"
                      f"+{B('am')}*{L['archive_member_arpu']}*{L['engagement_upsell']}*{L['drops_dau_share']}/0.3-{B('contra')}"),
            "a_ds": f"={B('Ra')}",
        }
        for k, label, fmt in LBLOCK:
            put(ws, lr[k], M0 + m - 1, f[k], fmt=fmt)
    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 9
    for m in range(1, H + 1):
        ws.column_dimensions[mc(m)].width = 15
    ws.freeze_panes = "C4"
    return lr


def sheet_mvp(wb, lr):
    ws = wb.create_sheet("MVP_Budget")
    ws["A1"] = "Бюджет MVP Loyalty (разработка дек–фев + 3 мес. пилота мар–май 2027)"
    ws["A1"].font = TITLE
    for j, h in enumerate(["Статья", "Сум", "USD", "Доля"], 1):
        ws.cell(row=3, column=j, value=h)
    hdr(ws, 3, 4)
    end = A.LOYALTY["mvp_start"] + A.LOYALTY["mvp_months"] - 1
    items = [("c_build", "Разработка MVP (web-view в приложении, CRM, антифрод)"), ("c_opex", "Операционные (команда, фулфилмент, КЦ)"),
             ("c_mkt", "Маркетинг MVP"), ("c_cb", "Кэшбэк"), ("c_dd", "Daily Drops — цифровые призы"),
             ("c_inv", "Daily Drops — инвентарь"), ("c_ww", "Weekly Wins"), ("c_gr", "Monthly Grand (1 розыгрыш)")]
    r = 4
    for k, label in items:
        put(ws, r, 1, label)
        put(ws, r, 2, f"=SUM(Loyalty_Monthly!{mc(1)}{lr[k]}:{mc(end)}{lr[k]})", GREEN, NUM)
        put(ws, r, 3, f"=B{r}/{INP['FX']}", BLACK, '"$"#,##0')
        put(ws, r, 4, f"=B{r}/B{4+len(items)}", BLACK, PCT)
        r += 1
    put(ws, r, 1, "ИТОГО MVP", BOLD)
    put(ws, r, 2, f"=SUM(B4:B{r-1})", BOLD, NUM)
    put(ws, r, 3, f"=B{r}/{INP['FX']}", BOLD, '"$"#,##0')
    put(ws, r + 1, 1, "На одного участника MVP, сум")
    put(ws, r + 1, 2, f"=B{r}/{LI['mvp_participants']}", BLACK, NUM)
    put(ws, r + 2, 1, "Без разработки (стоимость самого теста), сум")
    put(ws, r + 2, 2, f"=B{r}-B4", BLACK, NUM)
    ws.column_dimensions["A"].width = 60
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 14


def sheet_thresholds(wb):
    ws = wb.create_sheet("Thresholds")
    ws["A1"] = "Порог участия в Weekly/Monthly: сравнение вариантов (стационарный режим масштаба, сум/мес)"
    ws["A1"].font = TITLE
    ws["A2"] = ("Кэшбэк, Daily Drops и OPEX одинаковы во всех вариантах и сюда не входят. "
                "Поведенческие параметры — гипотезы; O2 и O5 тестируются двумя ветками MVP.")
    heads = ["Вариант", "Описание", "Доля < 60 тыс., которые играют", "Вовлечённость < 60 тыс.", "Переходы на ≥ 60 тыс./мес",
             "Club+ Pass: доля", "Цена Pass, с НДС", "Платный вход: доля", "Цена входа (средн.), с НДС",
             "P(запрет/предписание)", "Коэф. переходов архивных",
             "Удержание", "Переходы", "Pass", "Платный вход", "Призы W+M", "ЧИСТЫЙ ЭФФЕКТ", "С учётом риска",
             "Играют всего", "Архивных играют", "Шанс выиграть за неделю", "Юр. риск", "PR-риск"]
    for j, h in enumerate(heads, 1):
        ws.cell(row=4, column=j, value=h)
    hdr(ws, 4, len(heads))
    ws.row_dimensions[4].height = 60
    pb_block = {"O1": 0.02, "O2": 0.03, "O3": 0.05, "O4": 0.40, "O5": 0.05}
    af = {"O1": 0.5, "O2": 0.8, "O3": 0.6, "O4": 0.7, "O5": 1.0}
    L = LI
    r = 5
    acc = f"(1-(1-{L['member_churn_base']})^12)/{L['member_churn_base']}"
    items_week = LI["ww_week_items"]
    for k, o in A.THRESHOLD_OPTIONS.items():
        put(ws, r, 1, k, BOLD)
        put(ws, r, 2, o["name"])
        vals = [o["play_share_below"], o["engagement_below"], o["upgrade_rate"], o["pass_adoption"], o["pass_price"],
                o["paid_entry_share"], o["paid_entry_price"], pb_block[k], af[k]]
        fm = [PCT, PCT, PCT2, PCT, NUM, PCT, NUM, PCT, NUM2]
        for j, (v, f_) in enumerate(zip(vals, fm), 3):
            put(ws, r, j, v, BLUE, f_)
        below, full = L["below60"], L["full"]
        f = {
            12: f"={full}*{L['member_churn_base']}*{L['churn_reduction']}*((1-{below})+{below}*D{r})*{acc}*{L['arpu_net']}*{L['margin']}",
            13: f"={full}*{below}*E{r}*{acc}*{L['upgrade_delta']}*{L['margin']}",
            14: f"={full}*{below}*F{r}*G{r}/(1+{INP['VAT']})*(1-{L['pass_value_cost']})",
            15: f"={full}*{below}*H{r}*I{r}/(1+{INP['VAT']})",
            16: f"=-({L['ww_week']}*52/12+{L['grand_month']})",
            17: f"=SUM(L{r}:P{r})",
            18: f"=Q{r}*(1-J{r})",
            19: f"={full}*((1-{below})+{below}*C{r})",
            20: f"=0.03+0.97*C{r}",
            21: f"=IF(S{r}>0,{items_week}/S{r},0)",
        }
        for j, v in f.items():
            fmt = PCT if j == 20 else ('0.0000%' if j == 21 else NUM)
            put(ws, r, j, v, BLACK, fmt)
        put(ws, r, 22, o["legal_risk"])
        put(ws, r, 23, o["pr_risk"])
        r += 1
    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 50
    for j in range(3, 24):
        ws.column_dimensions[CL(j)].width = 15
    ws.column_dimensions["V"].width = 36
    ws.freeze_panes = "C5"


# ============================================================================ Portfolio, AB, MC
def sheet_portfolio(wb, keys, brow, lr):
    ws = wb.create_sheet("Portfolio")
    ws["A1"] = "Портфель: ARPU архивной когорты (план: всё запущено) и вклад"
    ws["A1"].font = TITLE
    ws.cell(row=3, column=1, value="Месяц")
    for m in range(1, H + 1):
        ws.cell(row=3, column=M0 + m - 1, value=m)
    hdr(ws, 3, M0 + H - 1)
    rows = ["I1–I3: выручка (после перекрытия)", "I1–I3: Δ абонентов (после перекрытия)", "Loyalty: доп. выручка архивных",
            "Loyalty: Δ архивных абонентов", "ARPU когорты, сум", "ΔARPU когорты", "ΔARPU только I1–I3",
            "Вклад I1–I3", "Вклад Loyalty (вся программа)", "Вклад портфеля накопленный"]
    for i, t in enumerate(rows, 4):
        ws.cell(row=i, column=1, value=t)
    h = INP["h"]
    for m in range(1, H + 1):
        X = mc(m)
        cidx = M0 + m - 1
        rv = "+".join(f"Monthly_I!{X}{brow[(k, 'rev')]}" for k in keys)
        ds = "+".join(f"Monthly_I!{X}{brow[(k, 'dsubs')]}" for k in keys)
        cr = "+".join(f"(Monthly_I!{X}{brow[(k, 'contrib')]}+Monthly_I!{X}{brow[(k, 'cost')]})*(1-{h})-Monthly_I!{X}{brow[(k, 'cost')]}" for k in keys)
        put(ws, 4, cidx, f"=({rv})*(1-{h})", fmt=NUM)
        put(ws, 5, cidx, f"=({ds})*(1-{h})", fmt=NUM)
        put(ws, 6, cidx, f"=Loyalty_Monthly!{X}{lr['a_rev']}", GREEN, NUM)
        put(ws, 7, cidx, f"=Loyalty_Monthly!{X}{lr['a_ds']}", GREEN, NUM)
        put(ws, 8, cidx, f"=(Monthly_I!{X}$5+{X}4+{X}6)/(Monthly_I!{X}$4+{X}5+{X}7)", fmt=NUM)
        put(ws, 9, cidx, f"={X}8/{INP['ARPU0']}-1", fmt=PCT)
        put(ws, 10, cidx, f"=(Monthly_I!{X}$5+{X}4)/(Monthly_I!{X}$4+{X}5)/{INP['ARPU0']}-1", fmt=PCT)
        put(ws, 11, cidx, f"={cr}", fmt=NUM)
        put(ws, 12, cidx, f"=Loyalty_Monthly!{X}{lr['contrib']}", GREEN, NUM)
        prev = f"N({CL(cidx-1)}13)+" if m > 1 else ""
        put(ws, 13, cidx, f"={prev}{X}11+{X}12", fmt=NUM)
    s = 15
    ws.cell(row=s, column=1, value="Итоги").font = BOLD
    tot = [
        ("ΔARPU когорты M12", f"={mc(12)}9", PCT), ("ΔARPU когорты M24", f"={mc(24)}9", PCT),
        ("ΔARPU когорты M36", f"={mc(36)}9", PCT), ("из них I1–I3 на M24", f"={mc(24)}10", PCT),
        ("ARPU когорты M24, сум", f"={mc(24)}8", NUM),
        ("Доп. выручка архивной когорты Y1", f"=SUM({mc(1)}4:{mc(12)}4)+SUM({mc(1)}6:{mc(12)}6)", NUM),
        ("Доп. выручка архивной когорты Y2", f"=SUM({mc(13)}4:{mc(24)}4)+SUM({mc(13)}6:{mc(24)}6)", NUM),
        ("Доп. выручка архивной когорты Y3", f"=SUM({mc(25)}4:{mc(36)}4)+SUM({mc(25)}6:{mc(36)}6)", NUM),
        ("Вклад портфеля 24 мес.", f"={mc(24)}13", NUM), ("Вклад портфеля 36 мес.", f"={mc(36)}13", NUM),
    ]
    for i, (lab, f, fm) in enumerate(tot, s + 1):
        ws.cell(row=i, column=1, value=lab)
        put(ws, i, 3, f, BOLD, fm)
    ch = LineChart()
    ch.title = "ΔARPU архивной когорты (план)"
    ch.y_axis.number_format = "0%"
    ch.x_axis.title = "Месяц"
    data = Reference(ws, min_col=1, max_col=M0 + H - 1, min_row=9, max_row=10)
    ch.add_data(data, from_rows=True, titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=M0, max_col=M0 + H - 1, min_row=3))
    ch.height, ch.width = 8, 24
    ws.add_chart(ch, f"A{s + len(tot) + 3}")
    ws.column_dimensions["A"].width = 44
    for m in range(1, H + 1):
        ws.column_dimensions[mc(m)].width = 15
    ws.freeze_panes = "C4"


def sheet_ab(wb):
    ws = wb.create_sheet("AB_Calculator")
    ws["A1"] = "Калькулятор выборок (двусторонний тест, α и мощность — ниже)"
    ws["A1"].font = TITLE
    items = [("Уровень значимости α", 0.05, PCT), ("Мощность", 0.80, PCT),
             ("CV ARPU (σ/μ)", 1.2, "0.00"), ("MDE ARPU, отн.", 0.02, PCT),
             ("Пропорция в контроле p1 (напр. отток)", 0.012, PCT2), ("Пропорция в тесте p2", 0.010, PCT2)]
    for i, (lab, v, f) in enumerate(items, 3):
        ws.cell(row=i, column=1, value=lab)
        put(ws, i, 2, v, BLUE, f, YELLOW)
    ws["A10"] = "N на группу — ARPU"
    put(ws, 10, 2, "=ROUNDUP(2*(NORMSINV(1-B3/2)+NORMSINV(B4))^2*B5^2/B6^2,0)", BOLD, NUM)
    ws["A11"] = "N на группу — пропорции"
    put(ws, 11, 2, ("=ROUNDUP((NORMSINV(1-B3/2)*SQRT(2*((B7+B8)/2)*(1-(B7+B8)/2))"
                    "+NORMSINV(B4)*SQRT(B7*(1-B7)+B8*(1-B8)))^2/(B8-B7)^2,0)"), BOLD, NUM)
    ws.column_dimensions["A"].width = 48
    ws.column_dimensions["B"].width = 16


def sheet_mc(wb):
    ws = wb.create_sheet("MC_Results")
    ws["A1"] = "Монте-Карло (20 000 прогонов) — значения из v2/model/output/results_v2.json"
    ws["A1"].font = TITLE
    path = os.path.join(HERE, "output", "results_v2.json")
    if not os.path.exists(path):
        return
    res = json.load(open(path, encoding="utf-8"))
    heads = ["#", "Инициатива", "Вклад 24м P10", "P50", "P90", "ΔARPU P50", "P(запуска)", "Шанс успеха", "Шанс ≥70% плана"]
    for j, h in enumerate(heads, 1):
        ws.cell(row=3, column=j, value=h)
    hdr(ws, 3, len(heads))
    r = 4
    for k, x in res["initiatives"].items():
        vals = [k, x["short"], *x["contrib_24_bn"], x["arpu_uplift_eval"][1], x["exec_prob"], x["success_chance"],
                x["business_case_chance"]]
        fm = [None, None, NUM1, NUM1, NUM1, PCT2, PCT, PCT, PCT]
        for j, (v, f) in enumerate(zip(vals, fm), 1):
            put(ws, r, j, v, BLUE, f)
        r += 1
    l = res["loyalty"]
    vals = ["I4", "Loyalty (вклад 36м)", *l["contrib_36_bn"], l["archive_arpu_uplift_m24"][1], l["exec_prob"],
            l["success_chance"], l["business_case_chance"]]
    fm = [None, None, NUM1, NUM1, NUM1, PCT2, PCT, PCT, PCT]
    for j, (v, f) in enumerate(zip(vals, fm), 1):
        put(ws, r, j, v, BLUE, f)
    r += 2
    p = res["portfolio"]
    ws.cell(row=r, column=1, value="Портфель").font = BOLD
    r += 1
    for lab, v, f in [("ΔARPU M24 P10", p["arpu_uplift_m24"][0], PCT), ("ΔARPU M24 P50", p["arpu_uplift_m24"][1], PCT),
                      ("ΔARPU M24 P90", p["arpu_uplift_m24"][2], PCT), ("P(ΔARPU ≥ 8%)", p["p_uplift_m24_ge"]["0.08"], PCT),
                      ("P(ΔARPU ≥ 10%)", p["p_uplift_m24_ge"]["0.1"], PCT), ("Вклад 36м P50, млрд", p["contrib_36_bn"][1], NUM1),
                      ("Бюджет MVP P50, млрд", l["mvp_cost_bn"][1], NUM2),
                      ("Затраты Loyalty в мес. при масштабе P50, млрд", l["cost_month_scale_bn"][1], NUM2)]:
        ws.cell(row=r, column=1, value=lab)
        put(ws, r, 3, v, BLUE, f)
        r += 1
    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 34
    for j in range(3, 10):
        ws.column_dimensions[CL(j)].width = 15


def build():
    wb = Workbook()
    sheet_readme(wb)
    sheet_inputs(wb)
    keys, row, brow = sheets_initiatives(wb)
    sheet_loyalty_inputs(wb)
    sheet_prizes(wb)
    sheet_drops(wb)
    lr = sheet_loyalty_monthly(wb)
    sheet_mvp(wb, lr)
    sheet_thresholds(wb)
    sheet_portfolio(wb, keys, brow, lr)
    sheet_ab(wb)
    sheet_mc(wb)
    for sh in wb.worksheets:
        for rr in sh.iter_rows():
            for c in rr:
                f = c.font
                if f is None or f.name != FONT:
                    c.font = Font(name=FONT, bold=bool(f and f.bold), color=f.color if f else None,
                                  size=f.size if f and f.size else 11)
    wb.calculation.fullCalcOnLoad = True
    wb.save(OUT)
    print("saved", OUT)
    return keys, row, brow, lr


if __name__ == "__main__":
    build()
